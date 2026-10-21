"""Excel-отчёт: сводка по товарам, полная история цен и графики (нативные диаграммы Excel)."""
from __future__ import annotations

import io
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import PricePoint, Product, Subscription
from .texts import rule_text

RUB = '#,##0 "₽"'
PCT = '+0.0%;-0.0%;0%'
DATE = "DD.MM.YYYY HH:MM"
HEADER_FILL = PatternFill("solid", fgColor="4C1D95")
HEADER_FONT = Font(bold=True, color="FFFFFF")
ZEBRA = PatternFill("solid", fgColor="F5F3FF")
THIN = Border(bottom=Side(style="thin", color="DDD6FE"))
GREEN = PatternFill("solid", fgColor="DCFCE7")
RED = PatternFill("solid", fgColor="FEE2E2")
MAX_CHARTS = 20


def _local(ts: datetime | None, tz: ZoneInfo) -> datetime | None:
    # Excel не умеет часовые пояса — пишем местное время без tzinfo
    return ts.astimezone(tz).replace(tzinfo=None) if ts else None


def _header(ws: Worksheet, row: int, titles: list[str], widths: list[int]) -> None:
    for col, (title, width) in enumerate(zip(titles, widths, strict=True), start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[row].height = 32
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _title(ws: Worksheet, text: str, subtitle: str) -> None:
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14, color="4C1D95")
    ws["A2"] = subtitle
    ws["A2"].font = Font(color="64748B")


def _summary(ws: Worksheet, items: list[tuple[Product, Subscription]], now: datetime, tz: ZoneInfo) -> None:
    _title(ws, "Мониторинг цен — сводка", f"Сформировано {_local(now, tz):%d.%m.%Y %H:%M} · товаров: {len(items)}")
    titles = ["Площадка", "Товар", "Цена сейчас", "Цена по карте", "Старая цена", "Скидка", "Цена на старте",
              "Минимум", "Максимум", "Изменение с начала", "В наличии", "Остаток", "Проверено", "Уведомления"]
    _header(ws, 4, titles, [13, 48, 13, 13, 13, 9, 13, 12, 12, 13, 10, 9, 17, 30])
    row = 5
    for i, (p, sub) in enumerate(items):
        change = (p.price - p.first_price) / p.first_price if p.price and p.first_price else None
        discount = (p.old_price - p.price) / p.old_price if p.price and p.old_price else None
        values = [p.marketplace.label, p.title, p.price, p.card_price, p.old_price, discount, p.first_price,
                  p.min_price, p.max_price, change, "да" if p.in_stock else "нет", p.stock_qty,
                  _local(p.checked_at, tz), rule_text(sub)]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = THIN
            cell.alignment = Alignment(vertical="center", wrap_text=col == 2)
            if i % 2:
                cell.fill = ZEBRA
        ws.cell(row=row, column=2).hyperlink = p.url
        ws.cell(row=row, column=2).font = Font(color="1D4ED8", underline="single")
        for col in (3, 4, 5, 7, 8, 9):
            ws.cell(row=row, column=col).number_format = RUB
        ws.cell(row=row, column=6).number_format = "0%"
        ws.cell(row=row, column=10).number_format = PCT
        ws.cell(row=row, column=13).number_format = DATE
        if not p.in_stock:
            ws.cell(row=row, column=11).font = Font(color="DC2626", bold=True)
        row += 1

    if items:
        last = row - 1
        ws.auto_filter.ref = f"A4:{get_column_letter(len(titles))}{last}"
        # подешевело — зелёным, подорожало — красным
        ws.conditional_formatting.add(f"J5:J{last}", CellIsRule(operator="lessThan", formula=["0"], fill=GREEN))
        ws.conditional_formatting.add(f"J5:J{last}", CellIsRule(operator="greaterThan", formula=["0"], fill=RED))


def _history(ws: Worksheet, items: list[tuple[Product, Subscription]], histories: dict[int, list[PricePoint]],
             tz: ZoneInfo) -> None:
    _title(ws, "История изменений цен", "Строка появляется при каждом изменении цены или наличия")
    _header(ws, 4, ["Дата", "Площадка", "Товар", "Цена", "Цена по карте", "В наличии"], [17, 13, 52, 13, 13, 10])
    rows = [(pt, p) for p, _ in items for pt in histories.get(p.id, [])]
    rows.sort(key=lambda r: r[0].ts, reverse=True)
    for i, (pt, p) in enumerate(rows, start=5):
        values = [_local(pt.ts, tz), p.marketplace.label, p.title, pt.price, pt.card_price,
                  "да" if pt.in_stock else "нет"]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=i, column=col, value=value)
            cell.border = THIN
        ws.cell(row=i, column=1).number_format = DATE
        ws.cell(row=i, column=4).number_format = RUB
        ws.cell(row=i, column=5).number_format = RUB
    if rows:
        ws.auto_filter.ref = f"A4:F{len(rows) + 4}"


def _charts(wb: Workbook, items: list[tuple[Product, Subscription]], histories: dict[int, list[PricePoint]],
            tz: ZoneInfo) -> None:
    charted = [(p, [pt for pt in histories.get(p.id, []) if pt.price]) for p, _ in items]
    charted = [(p, pts) for p, pts in charted if len(pts) >= 2][:MAX_CHARTS]
    if not charted:
        return
    ws = wb.create_sheet("Графики")
    data = wb.create_sheet("Данные графиков")
    data.sheet_state = "hidden"
    _title(ws, "Динамика цен", "Товары, у которых цена менялась хотя бы раз")
    col = 1
    for n, (p, pts) in enumerate(charted):
        data.cell(row=1, column=col, value="Дата")
        data.cell(row=1, column=col + 1, value=p.title[:60])
        for r, pt in enumerate(pts, start=2):
            data.cell(row=r, column=col, value=_local(pt.ts, tz).strftime("%d.%m %H:%M"))
            data.cell(row=r, column=col + 1, value=pt.price)
        chart = LineChart()
        chart.title = f"{p.marketplace.label}: {p.title[:55]}"
        chart.y_axis.title = "₽"
        chart.y_axis.numFmt = "#,##0"
        chart.height, chart.width = 7.5, 17
        chart.legend = None
        chart.add_data(Reference(data, min_col=col + 1, min_row=1, max_row=len(pts) + 1), titles_from_data=True)
        chart.set_categories(Reference(data, min_col=col, min_row=2, max_row=len(pts) + 1))
        chart.series[0].graphicalProperties.line.solidFill = "6D28D9"
        chart.series[0].graphicalProperties.line.width = 28000
        chart.series[0].smooth = False
        ws.add_chart(chart, f"A{4 + n * 16}")
        col += 3


def build_report(items: list[tuple[Product, Subscription]], histories: dict[int, list[PricePoint]],
                 now: datetime, tz: ZoneInfo) -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = "Сводка"
    _summary(summary, items, now, tz)
    _history(wb.create_sheet("История цен"), items, histories, tz)
    _charts(wb, items, histories, tz)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
