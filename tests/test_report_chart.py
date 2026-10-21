import io

from openpyxl import load_workbook

from bot.chart import price_chart
from bot.links import parse_product_url
from bot.report import build_report

from conftest import TZ, snap


async def make_items(db, clock):
    ref = parse_product_url("https://www.ozon.ru/product/5343983969/")
    product = await db.add_product(ref, snap(1846, card_price=1681, old_price=20000), clock(), clock())
    await db.upsert_user(1, None, "Тест")
    await db.subscribe(1, product.id, baseline=1846)
    for price in (1790, 1690, 1750):
        clock.advance(hours=6)
        await db.save_check(product.id, snap(price, card_price=price - 160), clock(), clock())
    items = await db.user_products(1)
    return items, {p.id: await db.history(p.id) for p, _ in items}


async def test_excel_report(db, clock):
    items, histories = await make_items(db, clock)
    wb = load_workbook(io.BytesIO(build_report(items, histories, clock(), TZ)))
    assert wb.sheetnames == ["Сводка", "История цен", "Графики", "Данные графиков"]

    summary = wb["Сводка"]
    header = [c.value for c in summary[4]]
    row = {h: c.value for h, c in zip(header, summary[5], strict=True)}
    assert row["Площадка"] == "Ozon"
    assert row["Цена сейчас"] == 1750 and row["Минимум"] == 1690 and row["Максимум"] == 1846
    assert round(row["Изменение с начала"], 4) == round((1750 - 1846) / 1846, 4)
    assert summary["B5"].hyperlink.target == "https://www.ozon.ru/product/5343983969/"

    assert wb["История цен"].max_row == 4 + 4  # 4 изменения цены
    assert len(wb["Графики"]._charts) == 1
    assert wb["Данные графиков"].sheet_state == "hidden"


async def test_report_without_history_has_no_charts(db, clock):
    ref = parse_product_url("https://www.avito.ru/moskva/telefony/iphone_8224149642")
    product = await db.add_product(ref, snap(28990), clock(), clock())
    await db.upsert_user(1, None, "Тест")
    await db.subscribe(1, product.id, baseline=28990)
    items = await db.user_products(1)
    wb = load_workbook(io.BytesIO(build_report(items, {product.id: await db.history(product.id)}, clock(), TZ)))
    assert "Графики" not in wb.sheetnames


async def test_price_chart_png(db, clock):
    items, histories = await make_items(db, clock)
    product = items[0][0]
    png = price_chart(product, histories[product.id], clock(), TZ)
    assert png.startswith(b"\x89PNG") and len(png) > 10_000
    assert price_chart(product, [], clock(), TZ).startswith(b"\x89PNG")  # пустая история не падает
